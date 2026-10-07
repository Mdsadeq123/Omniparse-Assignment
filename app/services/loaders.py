from pathlib import Path

from langchain_core.documents import Document
from pypdf import PdfReader

from app.logging_config import get_logger

logger = get_logger("parsing")

def load_document(file_path: Path, doc_id: str, filename: str) -> list[Document]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(file_path, doc_id, filename)
    if suffix == ".docx":
        return _load_docx(file_path, doc_id, filename)
    if suffix in {".txt", ".md", ".markdown"}:
        return _load_text(file_path, doc_id, filename)
    raise ValueError(f"Unsupported document type: {suffix}")


def _load_pdf(file_path: Path, doc_id: str, filename: str) -> list[Document]:
    reader = PdfReader(str(file_path))
    logger.info(
        "=== [PARSING / OCR] Extracting text from %s using PyPDF (%s pages) ===",
        filename,
        len(reader.pages),
    )
    docs: list[Document] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        docs.append(
            Document(
                page_content=text,
                metadata={
                    "doc_id": doc_id,
                    "source": filename,
                    "type": "document",
                    "page": str(i + 1),
                },
            )
        )
    logger.info(
        "[PARSING / OCR] PyPDF extraction complete: %s pages with text, %s characters",
        len(docs),
        sum(len(doc.page_content) for doc in docs),
    )
    return docs


def _load_docx(file_path: Path, doc_id: str, filename: str) -> list[Document]:
    from docx import Document as DocxDocument

    logger.info("=== [PARSING / OCR] Extracting text from %s using python-docx ===", filename)
    doc = DocxDocument(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    if not paragraphs:
        return []
    full_text = "\n\n".join(paragraphs)
    logger.info(
        "[PARSING / OCR] DOCX extraction complete: %s paragraphs, %s characters",
        len(paragraphs),
        len(full_text),
    )
    return [
        Document(
            page_content=full_text,
            metadata={
                "doc_id": doc_id,
                "source": filename,
                "type": "document",
                "page": "1",
            },
        )
    ]


def _load_text(file_path: Path, doc_id: str, filename: str) -> list[Document]:
    logger.info("=== [PARSING / OCR] Reading text from %s as UTF-8 ===", filename)
    text = file_path.read_text(encoding="utf-8", errors="replace")
    if not text.strip():
        return []
    logger.info("[PARSING / OCR] Text extraction complete: %s characters", len(text))
    return [
        Document(
            page_content=text,
            metadata={
                "doc_id": doc_id,
                "source": filename,
                "type": "document",
                "page": "1",
            },
        )
    ]
